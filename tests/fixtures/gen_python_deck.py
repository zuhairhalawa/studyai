"""Builds a deck with python-pptx (independent of our own fixture builder) and
dumps what python-pptx itself reads back, so tests can use it as an oracle.

usage: python -I gen_python_deck.py out.pptx expected.json
"""
import json
import struct
import sys
import tempfile
import zlib
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

out_pptx, out_json = sys.argv[1], sys.argv[2]


def png_bytes(w=8, h=8):
    raw = b"".join(b"\x00" + bytes([40, 90, 200] * w) for _ in range(h))

    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


prs = Presentation()

# 1 — title slide
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Information Security Basics"
s.placeholders[1].text = "Lecture 3 – Cryptography & Key Exchange"

# 2 — bullets with levels + notes
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "The CIA triad"
tf = s.placeholders[1].text_frame
tf.text = "Confidentiality"
for text, level in [
    ("Only authorised parties can read data", 1),
    ("Integrity", 0),
    ("Data is not modified without detection", 1),
    ("Availability", 0),
]:
    p = tf.add_paragraph()
    p.text = text
    p.level = level
s.notes_slide.notes_text_frame.text = "Mention the Equifax breach.\nAsk who has seen a DDoS."

# 3 — table
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Symmetric vs asymmetric"
rows = [
    ["Property", "Symmetric", "Asymmetric"],
    ["Keys", "One shared key", "Public/private pair"],
    ["Speed", "Fast", "Slow"],
]
tbl = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.8), Inches(9), Inches(2)).table
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        tbl.cell(r, c).text = val

# 4 — picture + rich text box
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Network layout"
with tempfile.TemporaryDirectory() as d:
    p = os.path.join(d, "x.png")
    with open(p, "wb") as f:
        f.write(png_bytes())
    s.shapes.add_picture(p, Inches(1), Inches(2), Inches(3), Inches(2))
tb = s.shapes.add_textbox(Inches(5), Inches(2), Inches(4), Inches(1.5))
para = tb.text_frame.paragraphs[0]
r1 = para.add_run()
r1.text = "A firewall "
r1.font.bold = True
r2 = para.add_run()
r2.text = "filters traffic"
r2.font.italic = True
r3 = para.add_run()
r3.text = " (see RFC 2979)"
r3.hyperlink.address = "https://www.rfc-editor.org/rfc/rfc2979"

# 5 — chart
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Attacks per year"
cd = CategoryChartData()
cd.categories = ["2021", "2022", "2023"]
cd.add_series("Phishing", (120, 180, 260))
cd.add_series("Ransomware", (40, 75, 130))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), cd)

# 6 — group shape
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Defence in depth"
grp = s.shapes.add_group_shape()
a = grp.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
a.text_frame.text = "Perimeter firewall"
b = grp.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(1))
b.text_frame.text = "Host intrusion detection"

prs.save(out_pptx)


# ---- oracle: read the saved file back with python-pptx itself ----------------
def texts_of(shape, out):
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            texts_of(sub, out)
        return
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    out.append(t)


saved = Presentation(out_pptx)
expected = []
for i, slide in enumerate(saved.slides, start=1):
    texts, pics, tables, charts = [], 0, 0, []
    for shape in slide.shapes:
        texts_of(shape, texts)
        if shape.shape_type == 13:
            pics += 1
        if getattr(shape, "has_table", False) and shape.has_table:
            tables += 1
        if getattr(shape, "has_chart", False) and shape.has_chart:
            ch = shape.chart
            charts.append(
                {
                    "categories": list(ch.plots[0].categories),
                    "series": [{"name": sr.name, "values": list(sr.values)} for sr in ch.plots[0].series],
                }
            )
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    expected.append(
        {
            "index": i,
            "title": slide.shapes.title.text if slide.shapes.title is not None else "",
            "texts": texts,
            "pictures": pics,
            "tables": tables,
            "charts": charts,
            "notes": notes,
        }
    )

with open(out_json, "w", encoding="utf-8") as f:
    json.dump({"slideCount": len(expected), "slides": expected}, f, ensure_ascii=False)
